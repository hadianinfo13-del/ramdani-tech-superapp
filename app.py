import streamlit as st
import google.generativeai as genai
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pandas as pd
import numpy as np
from PIL import Image
import io
import time
import sqlite3
import os
from datetime import datetime

# Konfigurasi Halaman Light Theme
st.set_page_config(
    page_title="SPOT Asisten Guru - EdTech Platform",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ------------------------------------------
# KUSTOM CSS (DESAIN TEMA TERANG & KECE)
# ------------------------------------------
st.markdown("""
<style>
    /* Reset & Background Utama */
    .stApp {
        background-color: #F8FAFC !important;
        color: #0F172A !important;
        font-family: 'Inter', sans-serif;
    }
    
    /* Header & Teks */
    h1, h2, h3, h4, h5, h6 {
        color: #1E293B !important;
        font-weight: 700 !important;
    }
    
    /* Custom Card/Container */
    div[data-testid="stVerticalBlockBorderWrapper"] {
        background-color: #FFFFFF !important;
        border-radius: 12px !important;
        border: 1px solid #E2E8F0 !important;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05), 0 2px 4px -1px rgba(0, 0, 0, 0.03) !important;
        padding: 18px !important;
    }

    /* Sidebar Styling */
    section[data-testid="stSidebar"] {
        background-color: #FFFFFF !important;
        border-right: 1px solid #E2E8F0 !important;
    }
    
    /* Primary Buttons */
    .stButton > button[kind="primary"] {
        background: linear-gradient(135deg, #2563EB 0%, #1D4ED8 100%) !important;
        color: #FFFFFF !important;
        border: none !important;
        border-radius: 8px !important;
        font-weight: 600 !important;
        padding: 0.5rem 1rem !important;
        box-shadow: 0 2px 4px rgba(37, 99, 235, 0.2) !important;
        transition: all 0.2s ease-in-out !important;
    }
    .stButton > button[kind="primary"]:hover {
        background: linear-gradient(135deg, #1D4ED8 0%, #1E40AF 100%) !important;
        box-shadow: 0 4px 8px rgba(37, 99, 235, 0.3) !important;
    }
    
    /* Secondary Buttons */
    .stButton > button {
        border-radius: 8px !important;
        border: 1px solid #CBD5E1 !important;
        color: #334155 !important;
        background-color: #FFFFFF !important;
        font-weight: 500 !important;
    }
    .stButton > button:hover {
        border-color: #94A3B8 !important;
        background-color: #F1F5F9 !important;
    }

    /* Custom Tabs */
    .stTabs [data-baseweb="tab-list"] {
        gap: 8px;
        border-bottom: 2px solid #E2E8F0;
    }
    .stTabs [data-baseweb="tab"] {
        height: 48px;
        border-radius: 8px 8px 0 0;
        padding-left: 16px;
        padding-right: 16px;
        color: #64748B;
        font-weight: 600;
        background-color: transparent;
    }
    .stTabs [aria-selected="true"] {
        color: #2563EB !important;
        border-bottom: 3px solid #2563EB !important;
        background-color: #EFF6FF !important;
    }

    /* Input Fields */
    .stTextInput > div > div > input, 
    .stTextArea > div > div > textarea,
    .stSelectbox > div > div > div {
        background-color: #FFFFFF !important;
        color: #0F172A !important;
        border: 1px solid #CBD5E1 !important;
        border-radius: 8px !important;
    }
    .stTextInput > div > div > input:focus, 
    .stTextArea > div > div > textarea:focus {
        border-color: #2563EB !important;
        box-shadow: 0 0 0 2px rgba(37, 99, 235, 0.2) !important;
    }

    /* Dataframe / Table */
    .stDataFrame {
        border: 1px solid #E2E8F0 !important;
        border-radius: 8px !important;
        overflow: hidden !important;
    }
    
    /* Metrics Box */
    div[data-testid="stMetric"] {
        background-color: #F8FAFC !important;
        border: 1px solid #E2E8F0 !important;
        padding: 12px 16px !important;
        border-radius: 10px !important;
    }
</style>
""", unsafe_allow_html=True)

# AMBIL API KEY DARI SECRETS
API_KEY = os.environ.get("GEMINI_API_KEY") 
if API_KEY:
    genai.configure(api_key=API_KEY)
else:
    st.error("⚠️ GEMINI_API_KEY belum diatur di Secrets Streamlit Cloud!")

# ------------------------------------------
# DATABASE SYSTEM
# ------------------------------------------
@st.cache_resource
def init_db():
    conn = sqlite3.connect('sekolah.db')
    cursor = conn.cursor()
    
    cursor.execute('CREATE TABLE IF NOT EXISTS nilai_siswa (nama TEXT, harian INTEGER, ujian INTEGER)')
    cursor.execute('CREATE TABLE IF NOT EXISTS users (username TEXT, password TEXT, role TEXT, nama_asli TEXT)')
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS absensi_mengajar (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            tanggal TEXT,
            mapel TEXT,
            materi TEXT,
            hadir INTEGER,
            sakit INTEGER,
            izin INTEGER,
            alpa INTEGER,
            catatan TEXT
        )
    ''')
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS tugas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            judul_tugas TEXT,
            deskripsi TEXT,
            deadline TEXT
        )
    ''')
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS pengumpulan_tugas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            tugas_id INTEGER,
            nama_siswa TEXT,
            waktu_upload TEXT,
            nama_file TEXT,
            status TEXT
        )
    ''')
    
    cursor.execute("SELECT COUNT(*) FROM users")
    if cursor.fetchone()[0] == 0:
        users = [
            ('guru', 'guru123', 'guru', 'Bpk. Ramdani, M.Pd.'),
            ('siswa', 'siswa123', 'siswa', 'Budi Santoso'),
            ('kepsek', 'kepsek123', 'kepsek', 'Dra. Hj. Nurhayati'),
            ('ortu', 'ortu123', 'ortu', 'Orang Tua Budi')
        ]
        cursor.executemany("INSERT INTO users VALUES (?, ?, ?, ?)", users)
        
    cursor.execute("SELECT COUNT(*) FROM nilai_siswa")
    if cursor.fetchone()[0] == 0:
        cursor.executemany("INSERT INTO nilai_siswa VALUES (?, ?, ?)", [('Budi Santoso', 85, 78), ('Siti Rahma', 95, 92), ('Andi Wijaya', 65, 70)])
    
    conn.commit()
    conn.close()

init_db()

@st.cache_data(ttl=5)
def load_data_nilai():
    conn = sqlite3.connect('sekolah.db')
    df = pd.read_sql_query("SELECT nama as 'Nama Siswa', harian as 'Harian', ujian as 'Ujian' FROM nilai_siswa", conn)
    conn.close()
    return df

def simpan_data_nilai(df):
    conn = sqlite3.connect('sekolah.db')
    df_sql = df.rename(columns={'Nama Siswa': 'nama', 'Harian': 'harian', 'Ujian': 'ujian'})
    df_sql.to_sql('nilai_siswa', conn, if_exists='replace', index=False)
    conn.close()
    st.cache_data.clear()

# ------------------------------------------
# GEMINI AI INTEGRATION
# ------------------------------------------
def panggil_gemini(prompt, gambar=None):
    kandidat_model = [
        "gemini-3.8-flash",
        "gemini-3.5-flash",
        "gemini-2.5-flash",
        "gemini-1.5-flash"
    ]
    
    for nama_model in kandidat_model:
        try:
            m = genai.GenerativeModel(nama_model)
            res = m.generate_content([prompt, gambar] if gambar else prompt)
            return res.text
        except Exception:
            continue

    try:
        available_models = [m.name for m in genai.list_models() if 'generateContent' in m.supported_generation_methods]
        for nama_model in available_models:
            try:
                m = genai.GenerativeModel(nama_model)
                res = m.generate_content([prompt, gambar] if gambar else prompt)
                return res.text
            except Exception:
                continue
    except Exception as e:
        st.error(f"Gagal koneksi AI: {str(e)}")

    st.error("Koneksi AI gagal. Periksa API Key Anda.")
    return None

def buat_file_word_rapi(teks, judul):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(judul.upper())
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 58, 138)
    
    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Dokumen Resmi - SPOT Asisten Guru AI\n" + "─"*50)
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(9)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    for line in teks.split('\n'):
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line)
            run.font.name = 'Arial'
            run.font.size = Pt(10.5)
            if line.startswith('#') or line.startswith('**'):
                run.font.bold = True
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(29, 78, 216)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

# ------------------------------------------
# AUTENTIKASI USER
# ------------------------------------------
if 'logged_in' not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.role = None
    st.session_state.nama_user = None

def login(username, password):
    conn = sqlite3.connect('sekolah.db')
    cursor = conn.cursor()
    cursor.execute("SELECT role, nama_asli FROM users WHERE username=? AND password=?", (username, password))
    user = cursor.fetchone()
    conn.close()
    if user:
        st.session_state.logged_in = True
        st.session_state.role = user[0]
        st.session_state.nama_user = user[1]
        st.rerun()
    else:
        st.error("❌ Username atau Password salah!")

def logout():
    st.session_state.logged_in = False
    st.session_state.role = None
    st.session_state.nama_user = None
    st.rerun()

# ------------------------------------------
# 1. HALAMAN LOGIN (LIGHT & CLEAN)
# ------------------------------------------
if not st.session_state.logged_in:
    st.markdown("<br>", unsafe_allow_html=True)
    col1, col2, col3 = st.columns([1, 1.8, 1])
    with col2:
        with st.container(border=True):
            st.markdown("<h2 style='text-align: center; color: #2563EB;'>🎓 SPOT EdTech</h2>", unsafe_allow_html=True)
            st.markdown("<p style='text-align: center; color: #64748B;'>Sistem Pembelajaran Terpadu & Asisten AI</p>", unsafe_allow_html=True)
            st.divider()
            
            with st.form("form_login"):
                user_input = st.text_input("Username", placeholder="Masukkan username Anda")
                pass_input = st.text_input("Password", type="password", placeholder="••••••••")
                submit = st.form_submit_button("Masuk Ke Portal", type="primary", use_container_width=True)
                if submit:
                    login(user_input, pass_input)
            
            st.markdown("""
            <div style="background-color: #F1F5F9; padding: 12px; border-radius: 8px; margin-top: 15px; font-size: 13px; color: #475569;">
                <b>🔑 Akun Demo Sistem:</b><br>
                • Guru: <code>guru</code> / <code>guru123</code><br>
                • Siswa: <code>siswa</code> / <code>siswa123</code><br>
                • Kepsek: <code>kepsek</code> / <code>kepsek123</code><br>
                • Ortu: <code>ortu</code> / <code>ortu123</code>
            </div>
            """, unsafe_allow_html=True)

# ------------------------------------------
# 2. PORTAL GURU
# ------------------------------------------
elif st.session_state.role == "guru":
    st.sidebar.markdown(f"### 👨‍🏫 {st.session_state.nama_user}")
    st.sidebar.caption("Pengajar / Dosen - SPOT LMS")
    st.sidebar.markdown("---")
    st.sidebar.button("🚪 Logout / Keluar", on_click=logout, use_container_width=True)
    
    st.title("👨‍🏫 Portal Pengajar SPOT AI")
    st.markdown("<p style='color: #64748B; margin-top: -15px;'>Kelola perkuliahan, jurnal mengajar, dan tugas siswa dalam satu tempat.</p>", unsafe_allow_html=True)
    
    tab1, tab2, tab3, tab4, tab5 = st.tabs([
        "📋 Jurnal & Absensi", 
        "📤 Management Tugas", 
        "📄 Generator Modul AI", 
        "📸 Scan Koreksi AI", 
        "📊 Buku Nilai"
    ])

    with tab1:
        st.subheader("📝 Jurnal & Absensi Perkuliahan / Kelas")
        with st.container(border=True):
            with st.form("form_absensi"):
                col1, col2 = st.columns(2)
                tgl = col1.date_input("Tanggal", datetime.now())
                mapel_input = col2.text_input("Mata Pelajaran / Kuliah", "Biologi X")
                materi_input = st.text_input("Materi Pembahasan", "Sistem Pencernaan Manusia")
                
                st.markdown("**Rekap Kehadiran Siswa:**")
                c1, c2, c3, c4 = st.columns(4)
                h = c1.number_input("Hadir", min_value=0, value=30)
                s = c2.number_input("Sakit", min_value=0, value=1)
                i = c3.number_input("Izin", min_value=0, value=0)
                a = c4.number_input("Alpa", min_value=0, value=0)
                catatan_input = st.text_area("Catatan Perkuliahan / Kendala", "Siswa sangat antusias saat diskusi kelompok.")
                
                if st.form_submit_button("💾 Simpan Jurnal Mengajar", type="primary"):
                    conn = sqlite3.connect('sekolah.db')
                    c = conn.cursor()
                    c.execute("INSERT INTO absensi_mengajar (tanggal, mapel, materi, hadir, sakit, izin, alpa, catatan) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                              (str(tgl), mapel_input, materi_input, h, s, i, a, catatan_input))
                    conn.commit()
                    conn.close()
                    st.success("✅ Jurnal Mengajar Berhasil Disimpan!")

        st.subheader("📜 Riwayat Jurnal Mengajar")
        conn = sqlite3.connect('sekolah.db')
        df_absen = pd.read_sql_query("SELECT tanggal as Tanggal, mapel as Mapel, materi as 'Materi Pembahasan', hadir as Hadir, sakit as Sakit, izin as Izin, alpa as Alpa FROM absensi_mengajar ORDER BY id DESC", conn)
        conn.close()
        st.dataframe(df_absen, use_container_width=True)

    with tab2:
        st.subheader("📢 Buat Penugasan Baru")
        with st.container(border=True):
            with st.form("form_tugas"):
                j_tugas = st.text_input("Judul Tugas", "Tugas 1: Analisis Organ Pencernaan")
                d_tugas = st.text_area("Instruksi / Detail Tugas", "Buatlah rangkuman materi dalam bentuk PDF minimal 2 halaman.")
                dl_tugas = st.date_input("Deadline Pengumpulan")
                if st.form_submit_button("🚀 Terbitkan Tugas", type="primary"):
                    conn = sqlite3.connect('sekolah.db')
                    c = conn.cursor()
                    c.execute("INSERT INTO tugas (judul_tugas, deskripsi, deadline) VALUES (?, ?, ?)", (j_tugas, d_tugas, str(dl_tugas)))
                    conn.commit()
                    conn.close()
                    st.success("Tugas berhasil diterbitkan ke siswa!")

        st.subheader("📥 Daftar Tugas Terkumpul")
        conn = sqlite3.connect('sekolah.db')
        df_kumpul = pd.read_sql_query('''
            SELECT p.nama_siswa as 'Nama Siswa', t.judul_tugas as 'Judul Tugas', p.waktu_upload as 'Waktu Upload', p.nama_file as 'Nama Berkas', p.status as Status 
            FROM pengumpulan_tugas p 
            JOIN tugas t ON p.tugas_id = t.id 
            ORDER BY p.id DESC
        ''', conn)
        conn.close()
        st.dataframe(df_kumpul, use_container_width=True)

    with tab3:
        st.subheader("📄 Generator RPP / Modul Ajar AI")
        mapel_rpp = st.text_input("Topik Pembelajaran:", "Fotosintesis pada Tumbuhan")
        if st.button("Generate RPP AI", type="primary") and mapel_rpp:
            with st.spinner("Merakit RPP dengan AI..."):
                prompt = f"Buatkan RPP Kurikulum Merdeka ringkas, rapi, dan terstruktur untuk topik {mapel_rpp}."
                hasil = panggil_gemini(prompt)
                if hasil:
                    st.success("✨ RPP Berhasil Dibuat!")
                    with st.container(border=True):
                        st.markdown(hasil)
                        file_word = buat_file_word_rapi(hasil, f"RPP - {mapel_rpp}")
                        st.download_button("📄 Unduh Dokumen (.docx)", file_word, f"RPP_{mapel_rpp}.docx", use_container_width=True)

    with tab4:
        st.subheader("📸 Scan & Koreksi Tugas AI Vision")
        file_foto = st.file_uploader("Upload Foto Lembar Jawaban Siswa", type=["jpg", "png", "jpeg"])
        if file_foto:
            img = Image.open(file_foto)
            st.image(img, caption="Preview Jawaban Siswa", width=350)
            if st.button("Koreksi Lembar Jawaban Ini", type="primary"):
                with st.spinner("Menganalisis tulisan..."):
                    prompt = "Tolong transkrip tulisan tangan di gambar ini, berikan penilaian (1-100) serta analisis kesalahan secara mendalam."
                    hasil = panggil_gemini(prompt, gambar=img)
                    if hasil:
                        with st.container(border=True):
                            st.markdown(hasil)

    with tab5:
        st.subheader("📊 Buku Nilai Terintegrasi")
        df_edit = st.data_editor(load_data_nilai(), num_rows="dynamic", use_container_width=True)
        if st.button("💾 Simpan Perubahan Nilai", type="primary"):
            simpan_data_nilai(df_edit)
            st.success("Buku Nilai Berhasil Diperbarui!")

# ------------------------------------------
# 3. PORTAL SISWA
# ------------------------------------------
elif st.session_state.role == "siswa":
    st.sidebar.markdown(f"### 👨‍🎓 {st.session_state.nama_user}")
    st.sidebar.caption("Siswa / Mahasiswa")
    st.sidebar.markdown("---")
    st.sidebar.button("🚪 Logout / Keluar", on_click=logout, use_container_width=True)
    
    st.title("🎓 Portal Belajar SPOT - Siswa")
    tab_s1, tab_s2 = st.tabs(["📤 Upload & Pengumpulan Tugas", "🤖 AI Tutor Personal"])

    with tab_s1:
        st.subheader("📢 Daftar Tugas Aktif")
        conn = sqlite3.connect('sekolah.db')
        c = conn.cursor()
        c.execute("SELECT id, judul_tugas, deskripsi, deadline FROM tugas ORDER BY id DESC")
        daftar_tugas = c.fetchall()
        conn.close()

        if daftar_tugas:
            for t_id, j_tugas, d_tugas, dl_tugas in daftar_tugas:
                with st.container(border=True):
                    st.markdown(f"#### 📌 {j_tugas}")
                    st.markdown(f"**Deadline:** `{dl_tugas}`")
                    st.write(f"**Instruksi:** {d_tugas}")
                    
                    file_upload = st.file_uploader(f"Unggah Berkas Tugas", type=["pdf", "docx", "png", "jpg"], key=f"file_{t_id}")
                    if st.button("📤 Kirim Tugas", key=f"btn_{t_id}", type="primary"):
                        if file_upload:
                            waktu_sekarang = datetime.now().strftime("%Y-%m-%d %H:%M")
                            conn = sqlite3.connect('sekolah.db')
                            cur = conn.cursor()
                            cur.execute("INSERT INTO pengumpulan_tugas (tugas_id, nama_siswa, waktu_upload, nama_file, status) VALUES (?, ?, ?, ?, ?)",
                                        (t_id, st.session_state.nama_user, waktu_sekarang, file_upload.name, "Terkirim Tepat Waktu"))
                            conn.commit()
                            conn.close()
                            st.success(f"✅ File '{file_upload.name}' berhasil dikirim!")
                        else:
                            st.warning("Pilih file terlebih dahulu.")
        else:
            st.info("Belum ada tugas aktif dari pengajar.")

    with tab_s2:
        st.subheader("🤖 Asisten Belajar AI")
        tanya = st.text_input("Tanyakan materi yang ingin kamu pelajari:")
        if st.button("Tanya AI", type="primary") and tanya:
            with st.spinner("Mencari jawaban..."):
                prompt = f"Jelaskan materi berikut secara ramah dan mudah dipahami oleh siswa: {tanya}"
                hasil = panggil_gemini(prompt)
                if hasil:
                    with st.container(border=True):
                        st.markdown(hasil)

# ------------------------------------------
# 4. PORTAL KEPALA SEKOLAH
# ------------------------------------------
elif st.session_state.role == "kepsek":
    st.sidebar.markdown(f"### 👔 {st.session_state.nama_user}")
    st.sidebar.caption("Kepala Sekolah / Dekan")
    st.sidebar.markdown("---")
    st.sidebar.button("🚪 Logout / Keluar", on_click=logout, use_container_width=True)
    
    st.title("📊 Executive Dashboard Monitoring")
    
    st.subheader("📌 Monitoring Jurnal Mengajar Guru")
    conn = sqlite3.connect('sekolah.db')
    df_jurnal_k = pd.read_sql_query("SELECT tanggal as Tanggal, mapel as Mapel, materi as 'Materi Pembahasan', hadir as Hadir, sakit as Sakit, izin as Izin, alpa as Alpa FROM absensi_mengajar", conn)
    conn.close()
    st.dataframe(df_jurnal_k, use_container_width=True)

    st.divider()
    st.subheader("📈 Rekap Rata-Rata Nilai Akademik")
    conn = sqlite3.connect('sekolah.db')
    df_k = pd.read_sql_query("SELECT * FROM nilai_siswa", conn)
    conn.close()
    if not df_k.empty:
        col1, col2 = st.columns(2)
        col1.metric("Rata-rata Nilai Harian", f"{df_k['harian'].mean():.1f}")
        col2.metric("Rata-rata Nilai Ujian", f"{df_k['ujian'].mean():.1f}")
        st.bar_chart(df_k.set_index('nama')['harian'])

# ------------------------------------------
# 5. PORTAL ORANG TUA
# ------------------------------------------
elif st.session_state.role == "ortu":
    st.sidebar.markdown(f"### 👨‍👩‍👧 {st.session_state.nama_user}")
    st.sidebar.caption("Wali Murid / Orang Tua")
    st.sidebar.markdown("---")
    st.sidebar.button("🚪 Logout / Keluar", on_click=logout, use_container_width=True)
    
    st.title("👨‍👩‍👧 Portal Monitoring Perkembangan Anak")
    nama_anak = st.text_input("🔍 Masukkan Nama Lengkap Anak:", "Budi Santoso")
    
    if st.button("Cari Laporan Anak", type="primary"):
        conn = sqlite3.connect('sekolah.db')
        df_o = pd.read_sql_query(f"SELECT * FROM nilai_siswa WHERE nama='{nama_anak}'", conn)
        df_t = pd.read_sql_query(f"SELECT t.judul_tugas, p.waktu_upload, p.status FROM pengumpulan_tugas p JOIN tugas t ON p.tugas_id = t.id WHERE p.nama_siswa='{nama_anak}'", conn)
        conn.close()

        if not df_o.empty:
            st.success(f"Laporan Akademik: **{nama_anak}**")
            col1, col2 = st.columns(2)
            col1.metric("Nilai Tugas Harian", df_o.iloc[0]['harian'])
            col2.metric("Nilai Ujian", df_o.iloc[0]['ujian'])
            
            st.subheader("📌 Riwayat Pengumpulan Tugas Anak")
            if not df_t.empty:
                st.dataframe(df_t, use_container_width=True)
            else:
                st.info("Belum ada catatan tugas yang dikumpulkan.")
        else:
            st.error("Data siswa tidak ditemukan. Pastikan nama lengkap sesuai.")
